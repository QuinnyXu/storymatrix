"""Build a StoryDB workbook from schema/storydb.schema.json.

    python tools/make_template.py                      -> templates/故事数据库模板.xlsx (blank, Chinese headers)
    python tools/make_template.py --lang en            -> templates/StoryDB_template_en.xlsx
    python tools/make_template.py --data templates/示例_星河旅店.data.json --out templates/示例_星河旅店.xlsx

The same schema drives the dashboard, so headers, dropdowns and the field notes always match what the tool reads.
Needs openpyxl (pip install openpyxl). Refuses to overwrite an existing file unless --force is given.
"""
import argparse
import datetime as dt
import json
import os
import re
import sys

sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCHEMA = json.load(open(os.path.join(ROOT, "schema", "storydb.schema.json"), encoding="utf-8"))
SHEETS = {s["key"]: s for s in SCHEMA["sheets"]}

FILL_REQ = PatternFill("solid", fgColor="1C5CAB")
FILL_OPT = PatternFill("solid", fgColor="CDE2FB")
FILL_AUTO = PatternFill("solid", fgColor="E1E0D9")
FILL_AUTO_BODY = PatternFill("solid", fgColor="F3F2EE")
FONT_REQ = Font(bold=True, color="FFFFFF")
FONT_OPT = Font(bold=True, color="0D366B")
FONT_AUTO = Font(italic=True, color="52514E")
WRAP = Alignment(wrap_text=True, vertical="top")

TXT = {
    "zh": {
        "readme": "说明", "noid": "⚠ ID不存在", "required": "必填", "optional": "选填", "auto": "自动（公式，不要手填）",
        "ref": "填「{sheet}」表里的 ID", "refs": "填「{sheet}」表里的 ID，多个用顿号分隔", "type": "类型",
        "how": [
            "怎么用",
            "1. 每张表第一列是 ID（如 C01、L01、T01、CH01）。ID 一旦给出就不改、不复用；不要的行把「状态」改成「弃」，不要删行。",
            "2. 深蓝表头 = 必填；浅蓝 = 选填；灰色「（自动）」= 公式，按 ID 显示名字，不要手填。鼠标停在表头上可以看到这一列怎么填。",
            "3. 带下拉箭头的格子可以从下拉里选（人物ID、地点ID、状态……），也可以手打。",
            "4. 一格里写多个 ID 用顿号分隔：C01、C02。",
            "5. 可以随意加新列（仪表盘会原样显示在详情里），也可以调整列的顺序；不要改已有的列名。",
            "6. 用不到的表可以删掉；「人物」和「事件」是必需的。",
            "7. 「出场」表写每个人在每个事件里做什么、在哪里；只想记「谁在场」时，填「事件.涉及人物ID」就够了。",
        ],
        "sheets": "各表", "col_sheet": "表", "col_desc": "用途",
    },
    "en": {
        "readme": "README", "noid": "⚠ no such ID", "required": "required", "optional": "optional", "auto": "auto (formula, do not type)",
        "ref": "an ID from the {sheet} sheet", "refs": "IDs from the {sheet} sheet, separated by 、 or commas", "type": "type",
        "how": [
            "How to use",
            "1. The first column of each sheet is an ID (C01, L01, T01, CH01). Never change or reuse an ID; to drop a row set Status to 弃 instead of deleting it.",
            "2. Dark-blue header = required; light blue = optional; grey '(auto)' = formula that shows the name for an ID. Hover a header to see how to fill it.",
            "3. Cells with a dropdown can be picked from the list (character IDs, location IDs, status…) or typed.",
            "4. Several IDs in one cell: C01、C02 (or commas).",
            "5. Add any new columns you like (the dashboard shows them in the details) and reorder freely; do not rename existing headers.",
            "6. Delete sheets you don't need; Characters and Events are required.",
            "7. Appearances records what each character does in each event and where; if you only need who is there, Events → Character IDs is enough.",
        ],
        "sheets": "Sheets", "col_sheet": "Sheet", "col_desc": "Purpose",
    },
}


def layout(sheet):
    return [f for f in sheet["fields"]]


def col_of(sheet_key, field_key):
    for i, f in enumerate(layout(SHEETS[sheet_key])):
        if f["key"] == field_key:
            return get_column_letter(i + 1)
    raise KeyError(field_key)


def qname(sheet_key, lang):
    return "'" + SHEETS[sheet_key][lang] + "'"


def auto_formula(sheet_key, f, row, lang):
    kind, src = f["auto"].split(":")
    target, name_field = {"name": ("characters", "name"), "loc": ("locations", "name"), "event": ("events", "title")}[kind]
    cell = f"${col_of(sheet_key, src)}{row}"
    key = cell
    src_field = next(x for x in SHEETS[sheet_key]["fields"] if x["key"] == src)
    if src_field.get("type") == "refs":  # show the first ID's name
        norm = f'SUBSTITUTE(SUBSTITUTE({cell},",","、"),"，","、")'
        key = f'LEFT({norm},FIND("、",{norm}&"、")-1)'
    t = qname(target, lang)
    return (f'=IF({cell}="","",IFERROR(INDEX({t}!${col_of(target, name_field)}:${col_of(target, name_field)},'
            f'MATCH({key},{t}!$A:$A,0)),"{TXT[lang]["noid"]}"))')


def note_for(f, lang):
    T = TXT[lang]
    if f.get("auto"):
        return T["auto"]
    parts = [T["required"] if f.get("required") else T["optional"]]
    if f.get("type") in ("ref", "refs"):
        parts.append(T[f["type"]].format(sheet=SHEETS[f["ref"]][lang]))
    if f.get("desc"):
        parts.append(f["desc"][lang])
    if f.get("enum"):
        parts.append(" / ".join(f["enum"]))
    if f.get("type") == "date":
        parts.append(SCHEMA["conventions"]["dateGrammar"][lang])
    if f.get("type") == "status":
        parts.append(" · ".join(f"{v} = {SCHEMA['conventions']['statusLabels'][v][lang]}" for v in SCHEMA["conventions"]["statusValues"]))
    if f.get("aliases"):
        parts.append(("也认：" if lang == "zh" else "Also read as: ") + "、".join(f["aliases"]))
    return "\n".join(parts)


def to_cell(v):
    if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
        y, m, d = map(int, v.split("-"))
        return dt.datetime(y, m, d)
    return v


def write_header(ws, sheet, lang):
    for i, f in enumerate(layout(sheet), start=1):
        c = ws.cell(row=1, column=i, value=f[lang])
        if f.get("auto"):
            c.fill, c.font = FILL_AUTO, FONT_AUTO
        elif f.get("required"):
            c.fill, c.font = FILL_REQ, FONT_REQ
        else:
            c.fill, c.font = FILL_OPT, FONT_OPT
        c.alignment = Alignment(vertical="center")
        c.comment = Comment(note_for(f, lang), "StoryDB", width=320, height=160)
        ws.column_dimensions[get_column_letter(i)].width = f.get("width", 12) * 1.15
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(sheet['fields']))}1"
    ws.row_dimensions[1].height = 22


def add_validations(ws, sheet, lang, nrows):
    statuses = ",".join(SCHEMA["conventions"]["statusValues"])
    for i, f in enumerate(layout(sheet), start=1):
        if f.get("auto"):
            continue
        rng = f"{get_column_letter(i)}2:{get_column_letter(i)}{nrows + 1}"
        formula = None
        if f.get("type") == "ref" and not f.get("lenient"):
            formula = f"{qname(f['ref'], lang)}!$A$2:$A$3000"
        elif f.get("type") == "status":
            formula = f'"{statuses}"'
        elif f.get("enum"):
            formula = '"' + ",".join(f["enum"]) + '"'
        if not formula:
            continue
        # warning style: the author may still type a value that is not in the list
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, errorStyle="warning", showErrorMessage=True,
                            error=("不在列表里，确定吗？" if lang == "zh" else "Not in the list - keep it?"))
        dv.add(rng)
        ws.add_data_validation(dv)


def build(lang, data, nrows):
    T = TXT[lang]
    wb = Workbook()
    rd = wb.active
    rd.title = T["readme"]
    title = (data or {}).get("settings", {}).get("title") or ("故事数据库" if lang == "zh" else "Story database")
    lines = [title, f"StoryDB v{SCHEMA['version']}", "", SCHEMA["about"][lang], "", *T["how"], "",
             ("日期写法：" if lang == "zh" else "Dates: ") + SCHEMA["conventions"]["dateGrammar"][lang], "",
             ("状态：" if lang == "zh" else "Status: ") + " · ".join(f"{v} = {SCHEMA['conventions']['statusLabels'][v][lang]}" for v in SCHEMA["conventions"]["statusValues"]),
             "", T["sheets"]]
    for i, line in enumerate(lines, start=1):
        c = rd.cell(row=i, column=1, value=line)
        if i == 1:
            c.font = Font(bold=True, size=14)
        elif line in (T["how"][0], T["sheets"]):
            c.font = Font(bold=True)
    r = len(lines) + 1
    for s in SCHEMA["sheets"]:
        rd.cell(row=r, column=1, value=s[lang] + ("" if s.get("optional") else " *")).font = Font(bold=True)
        rd.cell(row=r, column=2, value=s["desc"][lang]).alignment = WRAP
        r += 1
    rd.column_dimensions["A"].width = 22
    rd.column_dimensions["B"].width = 90

    for s in SCHEMA["sheets"]:
        ws = wb.create_sheet(s[lang])
        write_header(ws, s, lang)
        rows = (data or {}).get(s["key"], [])
        if s.get("keyValue"):
            given = (data or {}).get("settings", {})
            for i, st in enumerate(s["settings"], start=2):
                ws.cell(row=i, column=1, value=st[lang])
                ws.cell(row=i, column=2, value=given.get(st["key"], st["default"]))
                ws.cell(row=i, column=3, value=st["desc"][lang]).font = Font(color="52514E")
            continue
        for ri, row in enumerate(rows, start=2):
            for ci, f in enumerate(layout(s), start=1):
                if not f.get("auto") and f["key"] in row:
                    c = ws.cell(row=ri, column=ci, value=to_cell(row[f["key"]]))
                    if isinstance(c.value, dt.datetime):
                        c.number_format = "yyyy-mm-dd"
        # extra (book-specific) columns in the data file: "extra": {"Header": value}
        extras = []
        for row in rows:
            for h in row.get("extra", {}):
                if h not in extras:
                    extras.append(h)
        for k, h in enumerate(extras):
            ci = len(s["fields"]) + 1 + k
            c = ws.cell(row=1, column=ci, value=h)
            c.fill, c.font = FILL_OPT, FONT_OPT
            ws.column_dimensions[get_column_letter(ci)].width = 16
            for ri, row in enumerate(rows, start=2):
                if h in row.get("extra", {}):
                    ws.cell(row=ri, column=ci, value=row["extra"][h])
        # formula helpers for every row the author is likely to use
        for ci, f in enumerate(layout(s), start=1):
            if f.get("auto"):
                for ri in range(2, nrows + 2):
                    c = ws.cell(row=ri, column=ci, value=auto_formula(s["key"], f, ri, lang))
                    c.fill, c.font = FILL_AUTO_BODY, FONT_AUTO
        add_validations(ws, s, lang, nrows)
    return wb


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--lang", choices=["zh", "en"], default="zh")
    ap.add_argument("--data")
    ap.add_argument("--out")
    ap.add_argument("--rows", type=int, default=300)
    ap.add_argument("--force", action="store_true")
    a = ap.parse_args()
    data = json.load(open(a.data, encoding="utf-8")) if a.data else None
    out = a.out or os.path.join(ROOT, "templates", "故事数据库模板.xlsx" if a.lang == "zh" else "StoryDB_template_en.xlsx")
    if os.path.exists(out) and not a.force:
        sys.exit(f"{out} exists; pass --force to overwrite")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    build(a.lang, data, a.rows).save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
